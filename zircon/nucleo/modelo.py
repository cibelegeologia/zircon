"""Modelo de planilha no formato que o Zircon lê, com as instruções numa segunda aba.

Os valores de exemplo são inventados (cinco grãos perto de 600 Ma), só para mostrar o formato.
"""
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..idiomas import N_, atual, tr
from .constantes import LAMBDA_235, LAMBDA_238

# idade (Ma), erro relativo 2σ de 207Pb/235U e de 206Pb/238U (%), rho
_GRAOS = ((598.0, 2.2, 1.6, 0.55), (603.5, 2.0, 1.5, 0.61), (600.8, 2.6, 1.8, 0.48),
          (596.2, 2.3, 1.7, 0.52), (605.1, 2.1, 1.6, 0.58))

INSTRUCOES = (
    N_("Esta planilha mostra o formato que o Zircon lê. Apague as linhas de exemplo e cole os seus dados no lugar."),
    N_("O Zircon lê só a primeira aba. Deixe os dados na primeira aba e as anotações em outras abas."),
    N_("A coluna Amostra diz de qual amostra é cada análise. Pode haver várias amostras no mesmo arquivo: o Zircon "
       "dá uma cor a cada uma, e você decide se calcula cada amostra separada ou se junta várias numa unidade "
       "(por exemplo, amostras da mesma camada). Também dá para adicionar outras planilhas ao mesmo projeto."),
    N_("O cabeçalho ocupa uma única linha, logo acima dos dados. Não mescle células nem ponha a unidade numa "
       "segunda linha: escreva a unidade no próprio nome da coluna, como \"± 2σ (%)\"."),
    N_("Uma linha por análise, sem linhas em branco no meio e sem notas embaixo da tabela."),
    N_("Cada coluna de erro fica logo à direita do valor a que se refere. O nome do erro diz se ele é 1σ ou 2σ "
       "e leva \"%\" se for relativo; sem \"%\", o erro é absoluto."),
    N_("rho é a correlação entre os erros de 207Pb/235U e de 206Pb/238U. Sem essa coluna, o Zircon usa rho = 0 e avisa."),
    N_("Para dados de Tera-Wasserburg, troque as razões por 238U/206Pb, seu erro, 207Pb/206Pb, seu erro e rho."),
    N_("Se você só tem idades (sem razões), use as colunas Amostra, Análise, Idade (Ma) e ± 2σ (Ma)."),
    N_("Colunas a mais (U, Th, Pb, Th/U, idades calculadas pelo laboratório) podem ficar, desde que à direita das "
       "razões. O Zircon ignora as que não reconhece."),
    N_("Grãos que você decidiu excluir saem desta aba (guarde-os em outra aba, com o motivo). Os discordantes não "
       "precisam sair: o Zircon os separa pelo limite de discordância que você escolhe no painel."),
    N_("Os números ficam como números, não como texto. Vírgula ou ponto decimal, tanto faz, mas o mesmo em toda a planilha."),
    N_("Ao importar, o Zircon mostra como reconheceu cada coluna. Confira antes de continuar."),
)


def _razoes(t_ma):
    return math.expm1(LAMBDA_235 * t_ma), math.expm1(LAMBDA_238 * t_ma)


def salvar_modelo(caminho):
    """Grava o modelo .xlsx em `caminho`, com os textos no idioma atual."""
    pt = atual().startswith("pt")
    cab = [tr("Amostra") if pt else "Sample", tr("Análise") if pt else "Analysis",
           "207Pb/235U", "± 2σ", "206Pb/238U", "± 2σ", "rho"]
    wb = Workbook()
    ws = wb.active
    ws.title = tr("Dados")[:31]
    ws.append(cab)
    for i, (t, e75, e68, rho) in enumerate(_GRAOS, 1):
        r75, r68 = _razoes(t)
        ws.append([tr("EXEMPLO"), f"{tr('EXEMPLO')}-{i:02d}", round(r75, 5), round(r75 * e75 / 100, 5),
                   round(r68, 6), round(r68 * e68 / 100, 6), rho])
    negrito, fundo = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font, c.fill = negrito, fundo
    for letra, largura in zip("ABCDEFG", (14, 14, 13, 11, 13, 11, 8)):
        ws.column_dimensions[letra].width = largura
    ws.freeze_panes = "A2"

    ins = wb.create_sheet(tr("Instruções")[:31])
    ins.column_dimensions["A"].width = 4
    ins.column_dimensions["B"].width = 100
    ins.append(["", tr("Como preparar a planilha para o Zircon")])
    ins["B1"].font = Font(bold=True, size=13)
    for i, texto in enumerate(INSTRUCOES, 1):
        ins.append([i, tr(texto)])
        ins.cell(row=i + 1, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ins.cell(row=i + 1, column=1).alignment = Alignment(vertical="top")
    wb.save(caminho)
